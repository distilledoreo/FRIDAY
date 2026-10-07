"""Document extraction for the desktop assistant API.

Call install_document_extraction(app, auth_dependencies) on the FastAPI app
that already serves the phone's search/voice endpoints.
"""
import base64
import io
import re
from urllib.parse import unquote

from fastapi import HTTPException, Request
from starlette.concurrency import run_in_threadpool

MAX_DOCUMENT_BYTES = 25 * 1024 * 1024
MAX_DOCUMENT_CHARS = 60_000
MAX_SCANNED_PAGES = 4
TEXT_SUFFIXES = {
    "txt", "md", "markdown", "csv", "tsv", "json", "jsonl", "xml", "yaml", "yml", "toml", "ini", "log", "html", "htm",
    "py", "kt", "kts", "java", "js", "ts", "tsx", "jsx", "c", "h", "cpp", "hpp", "cs", "go", "rs", "rb", "php", "swift",
    "sh", "bash", "zsh", "sql", "css", "scss", "gradle", "properties", "srt", "vtt", "tex", "rtf",
}


def _clip(text: str) -> tuple[str, bool]:
    text = re.sub(r"[ \t]+\n", "\n", text)
    text = re.sub(r"\n{3,}", "\n\n", text).strip()
    return (text[:MAX_DOCUMENT_CHARS], True) if len(text) > MAX_DOCUMENT_CHARS else (text, False)


def extract_document(data: bytes, name: str, mime: str) -> dict:
    suffix = name.rsplit(".", 1)[-1].lower() if "." in name else ""
    if data[:5] == b"%PDF-" or suffix == "pdf" or mime == "application/pdf":
        import pypdf
        import pypdfium2

        reader = pypdf.PdfReader(io.BytesIO(data))
        pages = [(page.extract_text() or "").strip() for page in reader.pages]
        text = "\n\n".join(f"[Page {i + 1}]\n{t}" for i, t in enumerate(pages) if t)
        result = {"kind": "pdf", "pages": len(pages), "images": []}
        if len(re.sub(r"\s+", "", text)) < 40 * max(1, min(len(pages), 3)):
            # No usable text layer (a scan): send page images for the vision model instead.
            pdf = pypdfium2.PdfDocument(data)
            for i in range(min(len(pdf), MAX_SCANNED_PAGES)):
                page = pdf[i]
                scale = 1280 / max(page.get_width(), page.get_height())
                image = page.render(scale=max(scale, 0.5)).to_pil().convert("RGB")
                out = io.BytesIO()
                image.save(out, format="JPEG", quality=85)
                result["images"].append(base64.b64encode(out.getvalue()).decode())
            result["note"] = (
                f"No text layer; sent {len(result['images'])} page image(s)"
                + (f" of {len(pdf)}" if len(pdf) > MAX_SCANNED_PAGES else "")
                + "."
            )
            text = ""
        result["text"], result["truncated"] = _clip(text)
        return result
    if data[:2] == b"PK" and (suffix == "docx" or "wordprocessingml" in mime):
        import docx

        document = docx.Document(io.BytesIO(data))
        parts = [p.text for p in document.paragraphs]
        for table in document.tables:
            for row in table.rows:
                parts.append(" | ".join(cell.text.strip() for cell in row.cells))
        text, truncated = _clip("\n".join(parts))
        return {"kind": "docx", "text": text, "truncated": truncated}
    if suffix in TEXT_SUFFIXES or mime.startswith("text/") or mime in ("application/json", "application/xml"):
        if b"\x00" in data[:4096]:
            raise HTTPException(415, "that file looks binary, not text")
        raw = data.decode("utf-8", errors="replace")
        if suffix in ("html", "htm") or mime == "text/html":
            raw = re.sub(r"(?is)<(script|style).*?</\1>", " ", raw)
            raw = re.sub(r"(?s)<[^>]+>", " ", raw)
        text, truncated = _clip(raw)
        return {"kind": "text", "text": text, "truncated": truncated}
    raise HTTPException(415, f"can't read {suffix or mime or 'this'} files yet; try PDF, Word, or a text file")


async def extract(request: Request):
    data = await request.body()
    if not data:
        raise HTTPException(400, "send the file as the request body")
    if len(data) > MAX_DOCUMENT_BYTES:
        raise HTTPException(413, "file larger than 25 MB")
    name = unquote(request.headers.get("x-filename", "document"))[:200]
    mime = request.headers.get("content-type", "").split(";")[0].strip().lower()
    try:
        result = await run_in_threadpool(extract_document, data, name, mime)
    except HTTPException:
        raise
    except Exception as e:
        raise HTTPException(422, f"couldn't read {name}: {e.__class__.__name__}")
    return {"name": name, "chars": len(result["text"]), **result}


def install_document_extraction(app, auth_dependencies):
    """Register /extract using the existing service's authentication dependency."""
    app.add_api_route("/extract", extract, methods=["POST"], dependencies=auth_dependencies)
